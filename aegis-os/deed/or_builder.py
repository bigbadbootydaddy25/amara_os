"""
OR Builder — populates WS_11-409-19_OR_2026-06-04.xlsx
from all five agent results.

Sheets:
  Index     — chain of title oldest to newest
  Vesting   — current owner, deed BK/PG, legal description
  Tax       — assessor record: account, ticket, values, class
  Wells     — OGWIS well data
  DEP       — DEP OOG permits and status
  Map       — PENDING (awaiting Scott Marek)
  Summary   — OR summary for the examiner
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment, Font, PatternFill, Border, Side
)
from openpyxl.utils import get_column_letter

from config import (
    OUTPUT_DIR, OR_FILENAME, PARCEL_ID, PARCEL_DISTRICT, PARCEL_COUNTY,
    PARCEL_ACRES, ASSIGNMENT_FROM, ASSIGNMENT_CLIENT, RUN_DATE,
)
from core.logger import get_logger

log = get_logger("OR_BUILDER")

# Styles
_HEADER_FILL  = PatternFill("solid", fgColor="1F3864")
_SECTION_FILL = PatternFill("solid", fgColor="2E75B6")
_ALT_FILL     = PatternFill("solid", fgColor="D6E4F7")
_GAP_FILL     = PatternFill("solid", fgColor="FF6B6B")
_PENDING_FILL = PatternFill("solid", fgColor="FFE066")
_WHITE_FONT   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_BOLD_FONT    = Font(name="Calibri", bold=True, size=11)
_NORMAL_FONT  = Font(name="Calibri", size=11)
_GAP_FONT     = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_THIN_BORDER  = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _hdr(ws, row: int, col: int, text: str, width_hint: int = 20) -> None:
    cell = ws.cell(row=row, column=col, value=text)
    cell.font   = _WHITE_FONT
    cell.fill   = _HEADER_FILL
    cell.border = _THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = width_hint


def _cell(ws, row: int, col: int, value, bold=False, fill=None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font   = _BOLD_FONT if bold else _NORMAL_FONT
    cell.border = _THIN_BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = fill


def build_or(results: dict) -> Path:
    """Build and save the OR Excel workbook from agent results."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / OR_FILENAME

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    _build_summary(wb, results)
    _build_index(wb, results.get("chain", {}))
    _build_vesting(wb, results.get("vest", {}))
    _build_tax(wb, results.get("tax", {}))
    _build_wells(wb, results.get("well", {}))
    _build_dep(wb, results.get("dep", {}))
    _build_map(wb)

    wb.save(str(out_path))
    log.info("OR saved: %s", out_path)
    return out_path


def _build_summary(wb, results: dict) -> None:
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = f"OWNERSHIP REPORT — PARCEL {PARCEL_ID}"
    t.font  = Font(name="Calibri", bold=True, size=16, color="1F3864")
    t.alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"].value = (
        f"{PARCEL_DISTRICT} District | {PARCEL_COUNTY} County, WV | "
        f"{PARCEL_ACRES} acres | Assignment: {ASSIGNMENT_FROM} / {ASSIGNMENT_CLIENT}"
    )
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(name="Calibri", size=12, italic=True)

    ws.merge_cells("A3:H3")
    ws["A3"].value = f"Run Date: {RUN_DATE} | Status: PRELIMINARY"
    ws["A3"].alignment = Alignment(horizontal="center")

    # Status grid
    headers = ["Agent", "Status", "Records Found", "Errors", "Notes"]
    for ci, h in enumerate(headers, 1):
        _hdr(ws, 5, ci, h)

    agent_order = ["chain", "vest", "tax", "well", "dep"]
    agent_labels = {
        "chain": "CHAIN — Clerk Index",
        "vest":  "VEST — Parcel Viewer",
        "tax":   "TAX — Assessor",
        "well":  "WELL — OGWIS",
        "dep":   "DEP — OOG",
    }

    for ri, key in enumerate(agent_order, 6):
        r = results.get(key, {})
        status = r.get("status", "NOT_RUN")
        fill = _ALT_FILL if status == "COMPLETE" else _GAP_FILL
        count = r.get("count", len(r.get("instruments", r.get("records", r.get("wells", [])))))
        errors = "; ".join(r.get("errors", [])) or "None"
        note = r.get("note", "")

        _cell(ws, ri, 1, agent_labels.get(key, key), bold=True)
        _cell(ws, ri, 2, status, fill=fill if status != "COMPLETE" else None)
        _cell(ws, ri, 3, count)
        _cell(ws, ri, 4, errors[:120])
        _cell(ws, ri, 5, note[:120])

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 45
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22


def _build_index(wb, chain: dict) -> None:
    ws = wb.create_sheet("Index")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:J1")
    ws["A1"].value = f"CHAIN OF TITLE — Parcel {PARCEL_ID} | {PARCEL_COUNTY} County, WV"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["#", "Type", "Grantor", "Grantee", "Book", "Page",
               "Date Recorded", "Date Instrument", "Description", "Notes/Gaps"]
    widths  = [5, 20, 28, 28, 8, 8, 16, 16, 35, 35]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        _hdr(ws, 3, ci, h, w)

    instruments = chain.get("instruments", [])
    if not instruments:
        ws.merge_cells("A4:J4")
        ws["A4"].value = (
            "NO INSTRUMENTS FOUND — Site blocked (403). "
            "Run from Mac to pull real data."
        )
        ws["A4"].fill = _GAP_FILL
        ws["A4"].font = _GAP_FONT
        ws["A4"].alignment = Alignment(horizontal="center")
    else:
        for ri, inst in enumerate(instruments, 4):
            alt = ri % 2 == 0
            fill = _ALT_FILL if alt else None
            _cell(ws, ri, 1,  inst.get("seq", ""),           fill=fill)
            _cell(ws, ri, 2,  inst.get("type", ""),          fill=fill)
            _cell(ws, ri, 3,  inst.get("grantor", ""),       fill=fill)
            _cell(ws, ri, 4,  inst.get("grantee", ""),       fill=fill)
            _cell(ws, ri, 5,  inst.get("book", ""),          fill=fill)
            _cell(ws, ri, 6,  inst.get("page", ""),          fill=fill)
            _cell(ws, ri, 7,  inst.get("date_recorded", ""), fill=fill)
            _cell(ws, ri, 8,  inst.get("date_instr", ""),    fill=fill)
            _cell(ws, ri, 9,  inst.get("description", ""),   fill=fill)
            _cell(ws, ri, 10, inst.get("notes", ""),         fill=fill)

    # Gaps
    gaps = chain.get("gaps", [])
    if gaps:
        row = 4 + max(len(instruments), 1) + 2
        ws.cell(row=row, column=1, value="GAPS / BREAKS IN TITLE:").font = _BOLD_FONT
        for i, gap in enumerate(gaps, row + 1):
            ws.merge_cells(f"A{i}:J{i}")
            ws[f"A{i}"].value = f"⚠ {gap}"
            ws[f"A{i}"].fill  = _GAP_FILL
            ws[f"A{i}"].font  = _GAP_FONT

    ws.freeze_panes = "A4"


def _build_vesting(wb, vest: dict) -> None:
    ws = wb.create_sheet("Vesting")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:C1")
    ws["A1"].value = f"VESTING — Parcel {PARCEL_ID} | mapwv.gov"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    rec = vest.get("record", {})
    fields = [
        ("Parcel ID",          rec.get("parcel_id",         "NOT FOUND")),
        ("Current Owner",      rec.get("owner_name",        "NOT FOUND")),
        ("Owner Address",      rec.get("owner_address",     "NOT FOUND")),
        ("Vesting Deed Book",  rec.get("deed_book",         "NOT FOUND")),
        ("Vesting Deed Page",  rec.get("deed_page",         "NOT FOUND")),
        ("Deed Date",          rec.get("deed_date",         "NOT FOUND")),
        ("Legal Description",  rec.get("legal_description", "NOT FOUND")),
        ("Acreage",            rec.get("acreage",           "NOT FOUND")),
        ("District",           rec.get("district",          "NOT FOUND")),
        ("Account Number",     rec.get("account_number",    "NOT FOUND")),
        ("Search Status",      vest.get("status",           "NOT_RUN")),
    ]

    for ri, (label, value) in enumerate(fields, 3):
        _cell(ws, ri, 1, label, bold=True)
        fill = _GAP_FILL if value == "NOT FOUND" else None
        _cell(ws, ri, 2, value, fill=fill)

    missing = vest.get("missing", [])
    if missing:
        row = 3 + len(fields) + 1
        ws.cell(row=row, column=1, value="MISSING FIELDS:").font = _BOLD_FONT
        ws.cell(row=row, column=2, value=", ".join(missing))

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 55


def _build_tax(wb, tax: dict) -> None:
    ws = wb.create_sheet("Tax")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    ws["A1"].value = f"TAX RECORDS — Parcel {PARCEL_ID} | Harrison County Assessor"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["District", "Account #", "Ticket #", "Owner",
               "Land Value", "Mineral Value", "Building Value",
               "Total Value", "Class", "Tax Year"]
    widths  = [14, 14, 14, 30, 14, 14, 14, 14, 10, 12]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        _hdr(ws, 3, ci, h, w)

    records = tax.get("records", [])
    if not records:
        ws.merge_cells("A4:J4")
        ws["A4"].value = (
            "NO TAX RECORDS FOUND — Site blocked (403). "
            "Run from Mac to pull real data."
        )
        ws["A4"].fill = _GAP_FILL
        ws["A4"].font = _GAP_FONT
        ws["A4"].alignment = Alignment(horizontal="center")
    else:
        for ri, r in enumerate(records, 4):
            alt = ri % 2 == 0
            fill = _ALT_FILL if alt else None
            _cell(ws, ri, 1,  r.get("district", ""),       fill=fill)
            _cell(ws, ri, 2,  r.get("account_number", ""), fill=fill)
            _cell(ws, ri, 3,  r.get("ticket_number", ""),  fill=fill)
            _cell(ws, ri, 4,  r.get("owner_name", ""),     fill=fill)
            _cell(ws, ri, 5,  r.get("land_value", ""),     fill=fill)
            _cell(ws, ri, 6,  r.get("mineral_value", ""),  fill=fill)
            _cell(ws, ri, 7,  r.get("building_value", ""), fill=fill)
            _cell(ws, ri, 8,  r.get("total_value", ""),    fill=fill)
            _cell(ws, ri, 9,  r.get("property_class", ""), fill=fill)
            _cell(ws, ri, 10, r.get("tax_year", ""),       fill=fill)

    ws.freeze_panes = "A4"


def _build_wells(wb, well: dict) -> None:
    ws = wb.create_sheet("Wells")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:K1")
    ws["A1"].value = f"WELL DATA — {PARCEL_COUNTY} County / {PARCEL_DISTRICT} District | WVGS OGWIS"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    note = well.get("note", "")
    ws.merge_cells("A2:K2")
    ws["A2"].value = note
    fill2 = _GAP_FILL if "FAILED" in note else (_PENDING_FILL if "UNKNOWN" in note else None)
    if fill2:
        ws["A2"].fill = fill2

    headers = ["API Number", "Well Name", "Operator", "Spud Date", "Completion",
               "Type", "Status", "Formation", "Total Depth", "Oil (BBL)", "Gas (MCF)"]
    widths  = [16, 24, 24, 14, 14, 14, 14, 18, 12, 12, 12]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        _hdr(ws, 4, ci, h, w)

    wells = well.get("wells", [])
    if not wells:
        ws.merge_cells("A5:K5")
        ws["A5"].value = note or "NO WELLS FOUND"
        ws["A5"].fill  = _PENDING_FILL if "UNKNOWN" in note else _ALT_FILL
        ws["A5"].alignment = Alignment(horizontal="center")
    else:
        for ri, w in enumerate(wells, 5):
            alt = ri % 2 == 0
            fill = _ALT_FILL if alt else None
            _cell(ws, ri, 1,  w.get("api_number", ""),         fill=fill)
            _cell(ws, ri, 2,  w.get("well_name", ""),          fill=fill)
            _cell(ws, ri, 3,  w.get("operator", ""),           fill=fill)
            _cell(ws, ri, 4,  w.get("spud_date", ""),          fill=fill)
            _cell(ws, ri, 5,  w.get("completion_date", ""),    fill=fill)
            _cell(ws, ri, 6,  w.get("well_type", ""),          fill=fill)
            _cell(ws, ri, 7,  w.get("well_status", ""),        fill=fill)
            _cell(ws, ri, 8,  w.get("formation", ""),          fill=fill)
            _cell(ws, ri, 9,  w.get("total_depth", ""),        fill=fill)
            _cell(ws, ri, 10, w.get("production_oil_bbl", ""), fill=fill)
            _cell(ws, ri, 11, w.get("production_gas_mcf", ""), fill=fill)

    ws.freeze_panes = "A5"


def _build_dep(wb, dep: dict) -> None:
    ws = wb.create_sheet("DEP")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    ws["A1"].value = f"DEP OOG RECORDS — {PARCEL_COUNTY} County / {PARCEL_DISTRICT} District | tagis.dep.wv.gov"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["API Number", "Well Name", "Operator", "Permit #",
               "Permit Date", "Permit Status", "Well Status", "Plug Date", "Violations"]
    widths  = [16, 24, 24, 14, 14, 16, 16, 14, 12]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        _hdr(ws, 3, ci, h, w)

    records = dep.get("records", [])
    if not records:
        ws.merge_cells("A4:I4")
        note = dep.get("note", "NO DEP RECORDS FOUND")
        ws["A4"].value = note
        ws["A4"].fill  = _PENDING_FILL if "UNKNOWN" in note else _ALT_FILL
        ws["A4"].alignment = Alignment(horizontal="center")
    else:
        for ri, r in enumerate(records, 4):
            alt = ri % 2 == 0
            fill = _ALT_FILL if alt else None
            _cell(ws, ri, 1, r.get("api_number", ""),    fill=fill)
            _cell(ws, ri, 2, r.get("well_name", ""),     fill=fill)
            _cell(ws, ri, 3, r.get("operator", ""),      fill=fill)
            _cell(ws, ri, 4, r.get("permit_number", ""), fill=fill)
            _cell(ws, ri, 5, r.get("permit_date", ""),   fill=fill)
            _cell(ws, ri, 6, r.get("permit_status", ""), fill=fill)
            _cell(ws, ri, 7, r.get("well_status", ""),   fill=fill)
            _cell(ws, ri, 8, r.get("plug_date", ""),     fill=fill)
            _cell(ws, ri, 9, r.get("violations", ""),    fill=fill)

    ws.freeze_panes = "A4"


def _build_map(wb) -> None:
    ws = wb.create_sheet("Map")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"].value = f"MAP SHEET — Parcel {PARCEL_ID}"
    ws["A1"].font  = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:D6")
    ws["A3"].value = (
        "PENDING — Awaiting Scott Marek\n\n"
        "Map and survey data to be added when plat is received.\n"
        f"Parcel: {PARCEL_ID} | District: {PARCEL_DISTRICT} | "
        f"County: {PARCEL_COUNTY} | Acres: {PARCEL_ACRES}"
    )
    ws["A3"].fill      = _PENDING_FILL
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A3"].font      = Font(name="Calibri", size=12, bold=True)

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 25
    ws.row_dimensions[3].height = 80
