"""EIA-860 annual generator data: downloads the year's zip, opens 3_1_Generator_*, and
reads the 'Retired and Canceled' tab. See app.config for the download URL template, which
needs verifying against the live eia.gov listing before first run (file/tab naming has
shifted slightly across years).
"""

import io
import zipfile
from datetime import date, datetime, timezone

import requests
from openpyxl import load_workbook

from app.config import settings

from .base import NormalizedRecord

SOURCE_NAME = "EIA-860 Retired and Canceled Generators"

# Maps our normalized field name to the column header(s) EIA has used for it.
COLUMN_ALIASES: dict[str, list[str]] = {
    "plant_name": ["Plant Name"],
    "plant_code": ["Plant Code"],
    "state": ["State"],
    "county": ["County"],
    "latitude": ["Latitude"],
    "longitude": ["Longitude"],
    "nameplate_capacity_mw": ["Nameplate Capacity (MW)"],
    "retirement_year": ["Retirement Year"],
    "technology": ["Technology"],
    "status": ["Status"],
}


def _find_header_row(ws) -> int:
    """EIA-860 sheets have one or two title rows above the real header row."""
    for row_idx in range(1, 6):
        values = [cell.value for cell in ws[row_idx]]
        non_empty = sum(1 for v in values if v not in (None, ""))
        if non_empty >= len(COLUMN_ALIASES) // 2:
            return row_idx
    raise RuntimeError("Could not locate the header row in the EIA-860 generator sheet.")


def _column_index(headers: list[str | None], aliases: list[str]) -> int | None:
    for alias in aliases:
        for i, header in enumerate(headers):
            if header and header.strip().lower() == alias.lower():
                return i
    return None


def fetch() -> list[NormalizedRecord]:
    retrieved_at = datetime.now(timezone.utc)
    zip_url = settings.eia860_zip_url_template.format(year=settings.eia860_year)

    resp = requests.get(zip_url, timeout=120)
    resp.raise_for_status()

    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        generator_file = next(
            (
                name
                for name in zf.namelist()
                if name.lower().startswith("3_1_generator") and name.lower().endswith((".xlsx", ".xls"))
            ),
            None,
        )
        if generator_file is None:
            raise RuntimeError(f"No 3_1_Generator file found in {zip_url}. Contents: {zf.namelist()}")
        workbook_bytes = zf.read(generator_file)

    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    sheet_name = next(
        (s for s in workbook.sheetnames if "retired" in s.lower() and "cancel" in s.lower()),
        None,
    )
    if sheet_name is None:
        raise RuntimeError(f"No 'Retired and Canceled' tab found. Sheets: {workbook.sheetnames}")
    sheet = workbook[sheet_name]

    header_row_idx = _find_header_row(sheet)
    headers = [str(c.value).strip() if c.value is not None else None for c in sheet[header_row_idx]]
    column_index = {key: _column_index(headers, aliases) for key, aliases in COLUMN_ALIASES.items()}

    records = []
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if all(v is None for v in row):
            continue
        record_data = {key: (row[idx] if idx is not None else None) for key, idx in column_index.items()}
        records.append(
            NormalizedRecord(
                data=record_data,
                source=SOURCE_NAME,
                source_url=zip_url,
                retrieved_at=retrieved_at,
                dataset_published_at=date(settings.eia860_year, 12, 31),
            )
        )
    return records
